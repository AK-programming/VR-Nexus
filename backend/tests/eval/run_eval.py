"""Score a tender analysis against the analyst-authored tracker for it.

Two input modes, because the useful comparison is old-versus-new and the old
runs only exist as shipped workbooks:

  --workbook PATH   score a requirements_and_matches.xlsx that the pipeline
                    already produced. No database, no API key, no network.
                    This is how the baseline was measured and how any
                    downloaded analysis pack can be checked.

  --tender-id UUID  score a tender sitting in the database, reading its
                    Requirement rows directly. Needs the app's normal
                    environment. Use this in the edit-prompt/re-run loop.

Examples
--------
    # baseline: what the pipeline shipped before the prompt rewrite
    python -m tests.eval.run_eval --workbook ~/Downloads/analysis/requirements_and_matches.xlsx \\
        --gold smeda_dashboard --verbose

    # after a re-run, from the database
    python -m tests.eval.run_eval --tender-id 974d5224-... --gold smeda_dashboard

    # record a run so later runs can be compared against it
    python -m tests.eval.run_eval --workbook out.xlsx --gold smeda_dashboard --save-baseline

Exit code is 1 only when --fail-under is given and recall falls below it, so
this is safe to run in CI as a recall gate without failing on precision, which
is the metric legitimately in tension with recall.
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from tests.eval.score import load_gold, report, score  # noqa: E402

BASELINE_PATH = Path(__file__).parent / "baselines.json"


def rows_from_workbook(path: Path) -> tuple[list[dict], dict]:
    """Read the Requirements sheet of a generated analysis workbook.

    Column labels are looked up by name rather than by position, because the
    workbook's columns are configurable per tender (and the sheet this has to
    read includes both the pre-rewrite 92-column layout and the current one).
    """
    try:
        import openpyxl
    except ImportError:
        raise SystemExit("openpyxl is needed to read a workbook: pip install openpyxl")

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if "Requirements" not in wb.sheetnames:
        raise SystemExit(f"{path} has no 'Requirements' sheet (found: {', '.join(wb.sheetnames)})")
    ws = wb["Requirements"]

    header = None
    rows: list[dict] = []
    for raw in ws.iter_rows(values_only=True):
        if header is None:
            header = {str(v).strip(): i for i, v in enumerate(raw) if v}
            continue

        def cell(*labels):
            for label in labels:
                if label in header:
                    value = raw[header[label]]
                    if value not in (None, ""):
                        return str(value).strip()
            return ""

        description = cell("Clause / Requirement Description")
        if not description:
            continue
        rows.append({
            "pages": cell("Page Number"),
            "section": cell("Section Name"),
            "reference": cell("Reference Number"),
            "description": description,
            "mandatory": cell("Mandatory (Yes/No)"),
            "marks": cell("Marks"),
            "coverage": cell("Coverage"),
        })

    # Marks total, to catch the band-summing bug: scoring bands are alternatives
    # for one criterion, and adding them reported a 100-mark tender as 222.
    marks_total = 0.0
    for r in rows:
        try:
            marks_total += float(r["marks"])
        except (TypeError, ValueError):
            pass
    return rows, {"marks_total": marks_total}


def rows_from_db(tender_id: str) -> tuple[list[dict], dict]:
    from app.core.database import SessionLocal
    from app.models.requirement import Requirement

    db = SessionLocal()
    try:
        found = (
            db.query(Requirement)
            .filter(Requirement.tender_id == tender_id)
            .order_by(Requirement.page_number, Requirement.created_at)
            .all()
        )
        rows = [{
            "pages": r.page_label or (str(r.page_number) if r.page_number is not None else ""),
            "section": r.section_name or "",
            "reference": r.clause_reference or "",
            "description": r.description or "",
            "mandatory": r.mandatory_raw or "",
            "marks": str(r.marks) if r.marks is not None else "",
            "coverage": "",
        } for r in found]
        marks_total = sum(float(r.marks) for r in found if r.marks is not None)
        return rows, {"marks_total": marks_total}
    finally:
        db.close()


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    source = parser.add_mutually_exclusive_group(required=True)
    source.add_argument("--workbook", type=Path, help="a generated requirements_and_matches.xlsx")
    source.add_argument("--tender-id", help="score a tender already in the database")
    parser.add_argument("--gold", required=True, help="gold set name, e.g. smeda_dashboard")
    parser.add_argument("--verbose", action="store_true", help="list missed gold rows and unmatched emitted rows")
    parser.add_argument("--json", action="store_true", help="emit machine-readable metrics instead of the scorecard")
    parser.add_argument("--save-baseline", action="store_true", help="record this run in baselines.json")
    parser.add_argument("--fail-under", type=float, metavar="RECALL",
                        help="exit 1 if recall is below this (0-1). The CI gate.")
    args = parser.parse_args(argv)

    gold = load_gold(args.gold)
    emitted, extra = (
        rows_from_workbook(args.workbook) if args.workbook else rows_from_db(args.tender_id)
    )
    result = score(emitted, gold["rows"], gold_name=gold["meta"].get("name", args.gold))

    metrics = {
        "gold": args.gold,
        "gold_rows": result.gold_rows,
        "emitted_rows": result.emitted_rows,
        "ratio": round(result.ratio, 3),
        "recall": round(result.recall, 4),
        "precision": round(result.precision, 4),
        "f1": round(result.f1, 4),
        "marks_total": extra.get("marks_total"),
        "marks_stated": gold["meta"].get("stated_technical_marks"),
    }

    if args.json:
        print(json.dumps(metrics, indent=2))
    else:
        print(report(result, gold["meta"], verbose=args.verbose))
        stated = metrics["marks_stated"]
        if stated:
            got = metrics["marks_total"]
            verdict = "matches" if abs(got - stated) < 0.5 else "DOES NOT MATCH"
            print(f"  marks: extracted {got:g} against a stated {stated:g} -> {verdict}\n")

    if args.save_baseline:
        store = json.loads(BASELINE_PATH.read_text()) if BASELINE_PATH.exists() else {}
        store[args.gold] = metrics
        BASELINE_PATH.write_text(json.dumps(store, indent=1) + "\n")
        print(f"  baseline recorded in {BASELINE_PATH.name}")
    elif BASELINE_PATH.exists() and not args.json:
        store = json.loads(BASELINE_PATH.read_text())
        prior = store.get(args.gold)
        if prior:
            print("  vs recorded baseline:")
            for key in ("emitted_rows", "recall", "precision", "f1"):
                was, now = prior.get(key), metrics.get(key)
                if isinstance(was, (int, float)) and isinstance(now, (int, float)):
                    arrow = "->" if now == was else ("up" if now > was else "down")
                    print(f"    {key:14s} {was} {arrow} {now}")
            print()

    if args.fail_under is not None and result.recall < args.fail_under:
        print(f"  FAIL: recall {result.recall:.1%} is below the {args.fail_under:.1%} gate")
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
