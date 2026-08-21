"""
Run this AFTER Postgres/pgvector/Redis are up (and after verify_task_4.py
passes, since Task 5 sits directly on top of Task 4's matches) to confirm
Task 5 (Output Assembly & Folder Creation, Stage 4) works end-to-end,
sub-task by sub-task:

  5.1.1  Main sheet generation
  5.1.2  Summary sheet
  5.1.3  Instructions sheet
  5.2.1  Root directory creation
  5.2.2  Populate root directory
  5.2.3  ZIP packaging for download
  5.3.1  Optional .docx/.pdf summary report generator

It creates one real Tender with three real Requirement rows (one with an
accepted evidence match, one with no match at all, one that doesn't need
evidence), a real dummy tender PDF and evidence file on disk, runs the
real run_output_assembly(), then inspects the actual .xlsx/.zip/.docx it
produced. It also drives the real /finalize and /download endpoints with
FastAPI's TestClient. Everything it creates - DB rows and files - is
deleted at the end.

Usage (from the backend/ folder, with the venv active):
    python verify_task_5.py
"""
import shutil
import sys
import uuid
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import inspect

from app.core.config import get_settings
from app.core.database import SessionLocal, engine
from app.models import Document, RequirementEvidenceMatch, Tender, User
from app.models.requirement import Requirement
from app.models.chunk import Chunk
from app.models.enums import (
    DocumentCategory,
    DocumentFileType,
    DocumentTrainingStatus,
    MatchReviewStatus,
    MatchType,
    TenderStatus,
    UserRole,
)
from app.services import output_assembly

EXPECTED_MAIN_HEADERS = [
    "Page Number", "Section", "Clause Ref", "Description", "Mandatory Flag",
    "Evaluation Impact", "Marks", "Evidence Required", "Matched File Path", "Remarks",
]


def main() -> int:
    print("=== Task 5 (5.1.1-3 / 5.2.1-3 / 5.3.1) verification ===\n")
    settings = get_settings()

    # --- 0. Schema check ---
    print("[0/8] Checking database schema...")
    inspector = inspect(engine)
    tender_columns = {c["name"] for c in inspector.get_columns("tenders")}
    missing = {"output_folder_path", "output_zip_path", "finalized_at"} - tender_columns
    if missing:
        print(f"  FAIL: tenders table is missing columns {missing}.")
        return 1
    print("  OK: tenders table has output_folder_path/output_zip_path/finalized_at.\n")

    db = SessionLocal()
    created_ids = {"documents": [], "chunks": [], "matches": []}
    tender = None
    user = None
    tmp_tender_file = None
    tmp_library_file = None

    # Purge debris from a previous crashed run of this script (matches
    # verify_task_4.py's leftover-cleanup approach).
    stale = db.query(Document).filter(Document.title.like("Test Evidence (%")).all()
    if stale:
        stale_ids = [d.id for d in stale]
        db.query(Chunk).filter(Chunk.document_id.in_(stale_ids)).delete(synchronize_session=False)
        db.query(Document).filter(Document.id.in_(stale_ids)).delete(synchronize_session=False)
        db.commit()
        print(f"[pre-cleanup] Removed {len(stale)} leftover document(s) from a previous failed run.\n")
    stale_tenders = db.query(Tender).filter(Tender.name == "verify_task_5 test tender").all()
    for t in stale_tenders:
        db.query(Requirement).filter(Requirement.tender_id == t.id).delete()
        db.query(Tender).filter(Tender.id == t.id).delete()
    if stale_tenders:
        db.commit()

    folder_path = None
    zip_path = None

    try:
        user = User(
            name="Verify Task 5",
            email=f"verify-task5-{uuid.uuid4().hex[:8]}@example.com",
            password_hash="not-a-real-hash",
            role=UserRole.USER,
        )
        db.add(user)
        db.flush()

        # --- 1. Real files on disk + a Tender + Requirements ---
        print("[1/8] Creating a Tender with an original file, and a matched evidence document...")
        Path(settings.TENDER_STORAGE_DIR).mkdir(parents=True, exist_ok=True)
        Path(settings.LIBRARY_STORAGE_DIR).mkdir(parents=True, exist_ok=True)

        tmp_tender_file = Path(settings.TENDER_STORAGE_DIR) / f"verify_task_5_{uuid.uuid4().hex[:8]}.pdf"
        tmp_tender_file.write_text("fake tender PDF content for verify_task_5")

        tmp_library_file = Path(settings.LIBRARY_STORAGE_DIR) / f"verify_task_5_{uuid.uuid4().hex[:8]}.docx"
        tmp_library_file.write_text("fake case study content for verify_task_5")

        document = Document(
            category=DocumentCategory.CASE_STUDY,
            title="Test Evidence (verify_task_5)",
            original_filename=tmp_library_file.name,
            file_path=str(tmp_library_file),
            file_type=DocumentFileType.DOCX,
            training_status=DocumentTrainingStatus.INDEXED,
            uploaded_by=user.id,
        )
        db.add(document)
        db.flush()
        created_ids["documents"].append(document.id)

        tender = Tender(
            name="verify_task_5 test tender",
            original_filename=tmp_tender_file.name,
            file_path=str(tmp_tender_file),
            uploaded_by=user.id,
        )
        db.add(tender)
        db.flush()

        req_matched = Requirement(
            tender_id=tender.id, page_number=1, section_name="Eligibility",
            clause_reference="C-1", description="Provide a case study as evidence.",
            is_mandatory=True, marks=10, evidence_required=True,
        )
        req_missing = Requirement(
            tender_id=tender.id, page_number=2, section_name="Eligibility",
            clause_reference="C-2", description="Provide a certificate no evidence exists for.",
            is_mandatory=True, marks=5, evidence_required=True,
        )
        req_no_evidence = Requirement(
            tender_id=tender.id, page_number=3, section_name="Technical",
            clause_reference="C-3", description="A requirement needing no evidence.",
            is_mandatory=False, marks=15, evidence_required=False,
        )
        db.add_all([req_matched, req_missing, req_no_evidence])
        db.flush()

        match = RequirementEvidenceMatch(
            requirement_id=req_matched.id,
            document_id=document.id,
            confidence_score=0.95,
            match_type=MatchType.AUTO,
            review_status=MatchReviewStatus.ACCEPTED,  # already reviewed, per TN-MTC-05
        )
        db.add(match)
        db.commit()
        created_ids["matches"].append(match.id)
        print(f"  OK: Tender {tender.id} created with 3 requirements "
              f"(1 matched+accepted, 1 with no evidence found, 1 not needing evidence).\n")

        # --- 2. Run Stage 4 for real ---
        print("[2/8] Running run_output_assembly()...")
        result = output_assembly.run_output_assembly(db, tender)
        db.refresh(tender)
        assert tender.status == TenderStatus.READY_FOR_REVIEW, f"expected READY_FOR_REVIEW, got {tender.status}"
        assert tender.output_folder_path and Path(tender.output_folder_path).is_dir()
        assert tender.output_zip_path and Path(tender.output_zip_path).is_file()
        folder_path = Path(tender.output_folder_path)
        zip_path = Path(tender.output_zip_path)
        print(f"  OK: status -> ready_for_review, folder + zip written to disk.\n")

        # --- 3. 5.1.1 Main Sheet ---
        print("[3/8] Testing 5.1.1 (Main Sheet)...")
        excel_path = next(folder_path.glob("*_Requirement_Tracker.xlsx"))
        wb = load_workbook(excel_path)
        assert wb.sheetnames[0] == "Main Sheet"
        main = wb["Main Sheet"]
        headers = [c.value for c in main[1]]
        assert headers == EXPECTED_MAIN_HEADERS, f"header mismatch: {headers}"

        rows = {row[3]: row for row in main.iter_rows(min_row=2, values_only=True)}  # keyed by Description
        matched_row = rows["Provide a case study as evidence."]
        assert matched_row[8] == str(tmp_library_file), f"expected matched file path, got {matched_row[8]}"
        missing_row = rows["Provide a certificate no evidence exists for."]
        assert missing_row[8] == "Missing Document", f"expected 'Missing Document', got {missing_row[8]}"
        assert "manual sourcing" in (missing_row[9] or "")
        no_evidence_row = rows["A requirement needing no evidence."]
        assert no_evidence_row[7] == "No" and not no_evidence_row[8]
        print("  OK: exact TN-OUT-01 column headers, matched path, 'Missing Document' flag, "
              "and blank-when-not-required all correct.\n")

        # --- 4. 5.1.2 Summary Sheet ---
        print("[4/8] Testing 5.1.2 (Summary Sheet)...")
        assert "Summary Sheet" in wb.sheetnames
        summary_values = [tuple(c.value for c in row) for row in wb["Summary Sheet"].iter_rows()]
        assert any(row[0] == "Total Requirements" and row[1] == 3 for row in summary_values)
        assert any(row[0] == "Eligibility" and row[1] == 2 and row[2] == 15 for row in summary_values), (
            "expected Eligibility section: 2 requirements, 15 marks"
        )
        print("  OK: requirement counts + marks breakdown by section present and correct.\n")

        # --- 5. 5.1.3 Instructions Sheet ---
        print("[5/8] Testing 5.1.3 (Instructions Sheet)...")
        assert "Instructions" in wb.sheetnames
        instructions_text = "\n".join(
            str(c.value) for row in wb["Instructions"].iter_rows() for c in row if c.value
        )
        assert "Missing Document" in instructions_text and "Matched File Path" in instructions_text
        print("  OK: Instructions sheet present and explains the report's fields.\n")

        # --- 6. 5.2.1 Root directory naming ---
        print("[6/8] Testing 5.2.1 (root directory naming)...")
        date_str = tender.created_at.strftime("%Y-%m-%d")
        expected_name = f"{output_assembly.sanitize_filename(tender.name)}_{date_str}"
        assert folder_path.name == expected_name, f"expected '{expected_name}', got '{folder_path.name}'"
        print(f"  OK: root directory is named '{folder_path.name}' ({{TenderName}}_{{Date}} per TN-OUT-02).\n")

        # --- 7. 5.2.2 Populate root directory ---
        print("[7/8] Testing 5.2.2 (populate root directory) + 5.2.3 (zip) + 5.3.1 (summary report)...")
        assert (folder_path / tmp_tender_file.name).is_file(), "original tender file not copied"
        assert excel_path.is_file()
        required_docs_dir = folder_path / "Required Documents"
        assert required_docs_dir.is_dir()
        assert (required_docs_dir / tmp_library_file.name).is_file(), "matched evidence file not copied"
        print("  OK: root directory contains the original tender, the tracker, and "
              "Required Documents/ with the matched evidence file.")

        with zipfile.ZipFile(zip_path) as zf:
            names = zf.namelist()
        assert any(tmp_tender_file.name in n for n in names)
        assert any("Required Documents" in n and tmp_library_file.name in n for n in names)
        print("  OK: 5.2.3 - .zip archive contains the full folder structure.")

        docx_candidates = list(folder_path.glob("*_Summary.docx"))
        assert docx_candidates, "expected a *_Summary.docx in the output folder"
        pdf_candidates = list(folder_path.glob("*_Summary.pdf"))
        if pdf_candidates:
            print(f"  OK: 5.3.1 - summary report generated as both .docx and .pdf.\n")
        else:
            print(f"  OK: 5.3.1 - summary report generated as .docx "
                  f"(.pdf skipped - LibreOffice not found on this host, which is fine, it's optional).\n")

        # --- 8. /finalize and /download via the real API ---
        print("[8/8] Testing /finalize (re-runs Stage 4) and /download via the real API...")
        try:
            from fastapi.testclient import TestClient
            from app.api.deps import get_current_user
            from app.main import app as fastapi_app

            fastapi_app.dependency_overrides[get_current_user] = lambda: user
            client = TestClient(fastapi_app)

            finalize_resp = client.post(f"/api/tenders/{tender.id}/finalize")
            assert finalize_resp.status_code == 200, finalize_resp.text
            db.refresh(tender)
            assert tender.status == TenderStatus.FINALIZED
            assert tender.finalized_at is not None
            print("  OK: POST /finalize -> status = finalized, finalized_at set.")

            download_resp = client.get(f"/api/tenders/{tender.id}/download")
            assert download_resp.status_code == 200, download_resp.text
            assert download_resp.headers["content-type"] == "application/zip"
            assert len(download_resp.content) > 0
            print("  OK: GET /download -> returns the .zip file.\n")

            del fastapi_app.dependency_overrides[get_current_user]
        except (ImportError, RuntimeError) as exc:
            print(f"  SKIPPED: install httpx (`pip install httpx`) to exercise the live "
                  f"API endpoints. ({exc})\n")

        print("Task 5 (5.1.1-3 / 5.2.1-3 / 5.3.1) verified successfully.")
        return 0

    except AssertionError as exc:
        print(f"\nFAIL: {exc}")
        return 1
    finally:
        db.rollback()
        for match_id in created_ids["matches"]:
            db.query(RequirementEvidenceMatch).filter(RequirementEvidenceMatch.id == match_id).delete()
        if tender is not None:
            db.query(Requirement).filter(Requirement.tender_id == tender.id).delete()
            db.query(Tender).filter(Tender.id == tender.id).delete()
        for chunk_id in created_ids["chunks"]:
            db.query(Chunk).filter(Chunk.id == chunk_id).delete()
        for doc_id in created_ids["documents"]:
            db.query(Document).filter(Document.id == doc_id).delete()
        if user is not None:
            db.query(User).filter(User.id == user.id).delete()
        db.commit()
        db.close()

        if tmp_tender_file and tmp_tender_file.exists():
            tmp_tender_file.unlink()
        if tmp_library_file and tmp_library_file.exists():
            tmp_library_file.unlink()
        if folder_path and folder_path.exists():
            shutil.rmtree(folder_path)
        if zip_path and zip_path.exists():
            zip_path.unlink()
        print("[cleanup] All test data and generated files removed.")


if __name__ == "__main__":
    sys.exit(main())
