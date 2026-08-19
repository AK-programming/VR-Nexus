"""File-format extraction: PDF / DOCX / PPTX / image → `RawDoc`.

Embedded images are written to disk under `storage/images/{document_id}/` and
only their **paths** travel onward (LIB-IDX-02). Image bytes are never embedded
or placed in chunk content — a vector of pixels is not comparable with a vector
of text, and storing base64 in a text column would poison retrieval.

Pages whose extracted text is suspiciously short are re-read with Tesseract OCR,
which is what makes a scanned certificate indexable.
"""
from __future__ import annotations

import io
import logging
import uuid
from pathlib import Path

from app.services.library import storage
from app.services.library.parsers.base import Page, RawDoc

logger = logging.getLogger(__name__)

# Below this many characters a page is treated as image-only and sent to OCR.
# Deliberately low: a title page legitimately holds very little text, and OCR is
# slow enough that we would rather miss a sparse page than OCR every page.
OCR_CHAR_THRESHOLD = 60

# Images smaller than this in either dimension are decorations (bullets, rules,
# logos repeated in a footer) and are not worth keeping.
MIN_IMAGE_DIMENSION = 80

IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".tiff", ".tif"}


class ExtractionError(RuntimeError):
    """Raised when a file cannot be read at all."""


# ---------------------------------------------------------------------------
# OCR
# ---------------------------------------------------------------------------

def _ocr_image_bytes(data: bytes) -> str:
    """Best-effort OCR. Returns "" if Tesseract or Pillow is unavailable."""
    try:
        import pytesseract
        from PIL import Image

        with Image.open(io.BytesIO(data)) as img:
            return pytesseract.image_to_string(img).strip()
    except Exception as exc:
        logger.warning("OCR failed: %s", exc)
        return ""


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

def _extract_pdf(path: Path, document_id: uuid.UUID, raw: RawDoc) -> None:
    import fitz  # PyMuPDF

    try:
        doc = fitz.open(path)
    except Exception as exc:
        raise ExtractionError(f"Could not open PDF: {exc}") from exc

    img_dir = storage.images_dir(document_id)

    with doc:
        for index, page in enumerate(doc, start=1):
            text = page.get_text("text").strip()
            image_paths: list[str] = []
            ocr_applied = False

            for img_index, img in enumerate(page.get_images(full=True)):
                xref = img[0]
                try:
                    pix = fitz.Pixmap(doc, xref)
                    if pix.width < MIN_IMAGE_DIMENSION or pix.height < MIN_IMAGE_DIMENSION:
                        continue
                    if pix.alpha or pix.n > 3:
                        # Normalise CMYK/alpha to plain RGB so the PNG is
                        # viewable by anything.
                        pix = fitz.Pixmap(fitz.csRGB, pix)
                    out = img_dir / f"p{index}_{img_index}.png"
                    pix.save(out)
                    image_paths.append(storage.relative_path(out))
                except Exception as exc:
                    raw.warnings.append(f"Page {index} image {img_index} skipped: {exc}")

            if len(text) < OCR_CHAR_THRESHOLD:
                try:
                    rendered = page.get_pixmap(dpi=200).tobytes("png")
                    ocr_text = _ocr_image_bytes(rendered)
                    if len(ocr_text) > len(text):
                        text = ocr_text
                        ocr_applied = True
                except Exception as exc:
                    raw.warnings.append(f"Page {index} OCR skipped: {exc}")

            raw.pages.append(
                Page(number=index, text=text, image_paths=image_paths, ocr_applied=ocr_applied)
            )


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------

# A DOCX has no pages until it is rendered, so blocks of this many paragraphs
# stand in for pages. The number only affects the granularity of page labels.
DOCX_PARAGRAPHS_PER_PAGE = 40


def _extract_docx(path: Path, document_id: uuid.UUID, raw: RawDoc) -> None:
    import docx

    try:
        document = docx.Document(str(path))
    except Exception as exc:
        raise ExtractionError(f"Could not open DOCX: {exc}") from exc

    lines: list[str] = []
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "").lower() if para.style is not None else ""
        # Re-mark headings as markdown so the section parsers can see structure
        # that would otherwise exist only as a style attribute.
        lines.append(f"## {text}" if style.startswith("heading") else text)

    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))

    img_dir = storage.images_dir(document_id)
    saved_images: list[str] = []
    for rel_id, rel in document.part.rels.items():
        if "image" not in rel.reltype:
            continue
        try:
            blob = rel.target_part.blob
            suffix = Path(rel.target_part.partname).suffix or ".png"
            out = img_dir / f"img_{len(saved_images)}{suffix}"
            out.write_bytes(blob)
            saved_images.append(storage.relative_path(out))
        except Exception as exc:
            raw.warnings.append(f"DOCX image {rel_id} skipped: {exc}")

    if not lines:
        raw.pages.append(Page(number=1, text="", image_paths=saved_images))
        return

    for page_index, start in enumerate(range(0, len(lines), DOCX_PARAGRAPHS_PER_PAGE)):
        block = lines[start : start + DOCX_PARAGRAPHS_PER_PAGE]
        raw.pages.append(
            Page(
                number=page_index + 1,
                text="\n".join(block),
                # Word does not tell us which paragraph an image sat beside, so
                # every image is attached to the first block rather than guessed
                # onto a page it may not belong to.
                image_paths=saved_images if page_index == 0 else [],
            )
        )


# ---------------------------------------------------------------------------
# PPTX
# ---------------------------------------------------------------------------

def _extract_pptx(path: Path, document_id: uuid.UUID, raw: RawDoc) -> None:
    from pptx import Presentation

    try:
        presentation = Presentation(str(path))
    except Exception as exc:
        raise ExtractionError(f"Could not open PPTX: {exc}") from exc

    img_dir = storage.images_dir(document_id)

    for slide_index, slide in enumerate(presentation.slides, start=1):
        lines: list[str] = []
        image_paths: list[str] = []

        for shape_index, shape in enumerate(slide.shapes):
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    # The title becomes a heading so a deck's structure reads
                    # the same way a document's does.
                    is_title = shape == slide.shapes.title
                    lines.append(f"## {text}" if is_title else text)

            if getattr(shape, "shape_type", None) is not None and hasattr(shape, "image"):
                try:
                    image = shape.image
                    out = img_dir / f"s{slide_index}_{shape_index}.{image.ext}"
                    out.write_bytes(image.blob)
                    image_paths.append(storage.relative_path(out))
                except Exception as exc:
                    raw.warnings.append(f"Slide {slide_index} image skipped: {exc}")

        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(f"Speaker notes: {notes}")

        raw.pages.append(
            Page(number=slide_index, text="\n".join(lines), image_paths=image_paths)
        )


# ---------------------------------------------------------------------------
# Standalone image
# ---------------------------------------------------------------------------

def _extract_image(path: Path, document_id: uuid.UUID, raw: RawDoc) -> None:
    """An uploaded image is its own evidence: OCR gives it searchable text, and
    the file itself is copied into the document's image directory so the chunk
    can link to it like any other."""
    img_dir = storage.images_dir(document_id)
    out = img_dir / f"page_1{path.suffix.lower()}"
    try:
        out.write_bytes(path.read_bytes())
        stored = [storage.relative_path(out)]
    except Exception as exc:
        raw.warnings.append(f"Could not copy source image: {exc}")
        stored = []

    text = _ocr_image_bytes(path.read_bytes())
    if not text:
        raw.warnings.append("No text recovered from image (OCR returned nothing).")

    raw.pages.append(Page(number=1, text=text, image_paths=stored, ocr_applied=bool(text)))


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------

def extract(path: str | Path, filename: str, document_id: uuid.UUID) -> RawDoc:
    """Extract text and images from a stored upload.

    `filename` is the original client name and decides the format; `path` is the
    UUID-named file actually on disk.
    """
    path = Path(path)
    if not path.exists():
        raise ExtractionError(f"Stored file is missing: {path}")

    raw = RawDoc(filename=filename)
    suffix = Path(filename).suffix.lower() or path.suffix.lower()

    if suffix == ".pdf":
        _extract_pdf(path, document_id, raw)
    elif suffix == ".docx":
        _extract_docx(path, document_id, raw)
    elif suffix == ".pptx":
        _extract_pptx(path, document_id, raw)
    elif suffix in IMAGE_EXTENSIONS:
        _extract_image(path, document_id, raw)
    else:
        raise ExtractionError(f"Unsupported file type: {suffix or '(none)'}")

    if not raw.text.strip():
        raw.warnings.append(
            "No text could be extracted. The file may be a scan without a usable "
            "OCR result, or may be empty."
        )

    return raw


# Pages read by `sample_text`. Enough to identify a document; few enough that the
# upload request stays responsive.
SAMPLE_PAGE_LIMIT = 3


def sample_text(data: bytes, filename: str, max_chars: int = 4000) -> str:
    """Cheap text sample for the upload-time duplicate check (LIB-UI-06).

    Runs in the request path, so it reads only the first few pages, extracts no
    images and never invokes OCR. Returns "" if the file cannot be sampled —
    the caller falls back to the hash check alone.
    """
    suffix = Path(filename).suffix.lower()
    buffer = io.BytesIO(data)

    try:
        if suffix == ".pdf":
            import fitz

            with fitz.open(stream=data, filetype="pdf") as doc:
                pages = [
                    doc[i].get_text("text")
                    for i in range(min(SAMPLE_PAGE_LIMIT, doc.page_count))
                ]
            return "\n".join(pages)[:max_chars].strip()

        if suffix == ".docx":
            import docx

            document = docx.Document(buffer)
            lines: list[str] = []
            total = 0
            for para in document.paragraphs:
                text = para.text.strip()
                if not text:
                    continue
                lines.append(text)
                total += len(text)
                if total >= max_chars:
                    break
            return "\n".join(lines)[:max_chars].strip()

        if suffix == ".pptx":
            from pptx import Presentation

            presentation = Presentation(buffer)
            lines = []
            for index, slide in enumerate(presentation.slides):
                if index >= SAMPLE_PAGE_LIMIT:
                    break
                for shape in slide.shapes:
                    if shape.has_text_frame and shape.text_frame.text.strip():
                        lines.append(shape.text_frame.text.strip())
            return "\n".join(lines)[:max_chars].strip()

    except Exception as exc:
        logger.warning("Text sampling failed for %s: %s", filename, exc)

    # Images are skipped deliberately: sampling one means OCR, which is far too
    # slow for a request handler.
    return ""

